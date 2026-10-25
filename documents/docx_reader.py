"""
documents/docx_reader.py
Extract text from DOCX files using python-docx.
"""

from __future__ import annotations

import io
from typing import BinaryIO, Union

try:
    from docx import Document
except ImportError:
    Document = None


def read_docx(file_input: Union[BinaryIO, bytes]) -> str:
    """Extract paragraphs and table text from a DOCX file."""
    if Document is None:
        raise RuntimeError("DOCX support is unavailable. Please install python-docx.")

    if isinstance(file_input, bytes):
        file_input = io.BytesIO(file_input)

    doc = Document(file_input)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]

    # Also extract any table cell text
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    return "\n\n".join(paragraphs)
